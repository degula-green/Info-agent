from __future__ import annotations

from io import BytesIO

from docx import Document

from app.application.report.render import render_report


def _template_bytes() -> bytes:
    document = Document()
    for line in (
        "## 周报模板（通用版）",
        "**汇报人：** [你的姓名]",
        "**周期：** 2026年9月28日 — 2026年10月4日",
        "### 一、本周工作完成情况",
        "| 序号 | 工作事项 | 完成情况 | 关键产出/数据 | 状态 |",
        "|---|---|---|---|---|",
        "| 1 | [事项名称] | [已完成/进行中/延期] | [具体结果] | ✅/🔄/⚠️ |",
        "### 二、下周工作计划",
        "| 优先级 | 工作事项 | 预期目标 | 计划完成时间 | 所需支持 |",
        "|---|---|---|---|---|",
        "| P0 | [最重要的事] | [可量化目标] | [日期] | [资源/协作] |",
        "### 三、思考与建议（可选）",
        "- **流程优化：** [发现的问题 + 建议]",
        "### 四、需领导决策/协调事项",
        "1. [具体事项，说明背景与选项]",
    ):
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_render_report_fills_slots_and_keeps_template() -> None:
    output = render_report(
        _template_bytes(),
        {
            "title": "上一周周报（2026-09-28 ~ 2026-10-04）",
            "field.汇报人": "张三",
            "field.周期": "2026-09-28 ~ 2026-10-04",
            "table.一_本周工作完成情况.0": ["1", "官网部署", "已完成", "已上线", "✅"],
            "table.二_下周工作计划.0": ["P0", "接口联调", "完成联调", "2026-10-09", "需要测试支持"],
            "bullet.三_思考与建议_可选.0": "把部署步骤整理成清单",
            "number.四_需领导决策_协调事项.0": "确认下周上线窗口",
        },
    )
    document = Document(BytesIO(output))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "## 上一周周报（2026-09-28 ~ 2026-10-04）" in texts
    assert "**汇报人：** 张三" in texts
    assert "**周期：** 2026-09-28 ~ 2026-10-04" in texts
    assert "| 1 | 官网部署 | 已完成 | 已上线 | ✅ |" in texts
    assert "| P0 | 接口联调 | 完成联调 | 2026-10-09 | 需要测试支持 |" in texts
    assert "- **流程优化：** 把部署步骤整理成清单" in texts
    assert "1. 确认下周上线窗口" in texts
    # Untouched template lines stay in place.
    assert "| 序号 | 工作事项 | 完成情况 | 关键产出/数据 | 状态 |" in texts
