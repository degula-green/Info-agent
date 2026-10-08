from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from docx import Document

from app.capabilities.report import (
    CAPABILITY_NAME,
    NO_TEMPLATE_MESSAGE,
    WeeklyReportCapability,
)
from app.kernel.execution_context import bind_execution_context
from app.kernel.execution_context import ExecutionContext
from app.providers.report import ReportDraft

SHANGHAI = ZoneInfo("Asia/Shanghai")
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

TEMPLATE_PARAGRAPHS = [
    "## 周报模板（通用版）",
    "**汇报人：** [你的姓名]",
    "**部门/项目：** [部门或项目名称]",
    "**周期：** 2026年9月28日",
    "### 一、本周工作完成情况",
    "| 序号 | 工作事项 | 完成情况 | 关键产出/数据 | 状态 |",
    "|---|---|---|---|---|",
    "| 1 | [事项名称] | [已完成/进行中/延期] | [具体结果] | ✅/🔄/⚠️ |",
    "| 2 | [事项名称] | [已完成/进行中/延期] | [具体结果] | ✅/🔄/⚠️ |",
    "### 二、本周遇到的问题与风险",
    "| 问题/风险 | 影响程度 | 已采取动作 | 需要支持 |",
    "|---|---|---|---|",
    "| [问题描述] | 高/中/低 | [已尝试的解决方案] | [需要谁、提供什么资源] |",
    "### 四、思考与建议（可选）",
    "- **流程优化：** [发现的问题 + 建议]",
]


def template_bytes() -> bytes:
    document = Document()
    for text in TEMPLATE_PARAGRAPHS:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakeRAG:
    def __init__(self, resource_ids=("r1",), chunks=()):
        self.resource_ids = list(resource_ids)
        self.chunks = list(chunks)
        self.source_bodies: list[dict] = []
        self.export_bodies: list[dict] = []

    def search_sources(self, body, **kwargs):
        self.source_bodies.append(body)
        return {"resource_ids": list(self.resource_ids)}

    def export_scope(self, body, **kwargs):
        """The complete-scope primitive person_scope now pages through."""

        self.export_bodies.append(body)
        wanted = set(body.get("resource_types") or ["message", "attachment"])
        items = []
        for chunk in self.chunks:
            if chunk[1] not in wanted:
                continue
            items.append(
                {
                    "chunk_id": chunk[4],
                    "resource_id": chunk[0],
                    "resource_type": chunk[1],
                    "title": chunk[2],
                    "sender": {"name": chunk[3]},
                    "conversation": {"id": "c1", "name": "项目群"},
                    "sent_at": "2026-09-30T02:00:00Z",
                    "text": chunk[5],
                }
            )
        return {"items": items, "next_offset": len(items), "has_more": False}

    def search_content(self, body, **kwargs):
        return {
            "items": [
                {
                    "resource_id": chunk[0],
                    "resource_type": chunk[1],
                    "title": chunk[2],
                    "sender": {"name": chunk[3]},
                    "conversation": {"id": "c1", "name": "项目群", "type": "group", "platform": "wechat"},
                    "sent_at": "2026-09-30T02:00:00Z",
                    "best_score": 1.0,
                    "chunks": [
                        {
                            "chunk_id": chunk[4],
                            "text": chunk[5],
                            "score": 1.0,
                        }
                    ],
                }
                for chunk in self.chunks
            ]
        }


class FakeKnowledge:
    def __init__(self, *, matches=None, attachments=None, template=None, subject="张三"):
        self.subject = subject
        self.matches = matches if matches is not None else [
            {
                "person_key": "p1",
                "display_name": "张三",
                "identity_ids": ["id-1"],
                "private_conversation_ids": ["conv-1"],
            }
        ]
        self.attachments = attachments if attachments is not None else []
        self.template = template if template is not None else template_bytes()
        self.opened: list[str] = []

    def resolve_person(self, *, owner_user_id, name, request_id="", trace_id=""):
        return {"subject_name": self.subject, "matches": self.matches}

    def search_attachments(self, *, owner_user_id, name, limit=50, request_id="", trace_id=""):
        return {"items": self.attachments}

    def open_attachment(self, *, owner_user_id, attachment_id, request_id="", trace_id=""):
        self.opened.append(attachment_id)
        return self.template


class FakeProvider:
    def __init__(self, values):
        self.values = values
        self.calls: list[dict] = []

    def generate(self, **kwargs):
        self.calls.append(kwargs)
        return ReportDraft(values=self.values, model_calls=1)


class FakeAttachments:
    def __init__(self, uploaded=None, blob=None):
        self.uploaded = uploaded or {}
        self.blob = blob if blob is not None else template_bytes()
        self.saved: list[dict] = []

    def metadata(self, attachment_id):
        return self.uploaded.get(attachment_id)

    def read(self, attachment_id):
        return self.blob

    def save(self, *, owner_id, file_name, mime_type, data):
        record = {
            "attachment_id": f"artifact-{len(self.saved) + 1}",
            "owner_id": owner_id,
            "file_name": file_name,
            "mime_type": mime_type,
            "size_bytes": len(data),
            "expires_at": "2026-10-08T05:00:00+00:00",
            "data": data,
        }
        self.saved.append(record)
        return record


def _context(source_ref: dict | None = None) -> ExecutionContext:
    return ExecutionContext(
        task_id="t1",
        plan_id="p1",
        step_id="s1",
        owner_user_id="u1",
        organization_id=None,
        request_id="req-1",
        trace_id="trace-1",
        source_type="chat",
        source_ref=dict(source_ref or {}),
    )


def _run(capability, arguments, source_ref: dict | None = None):
    with bind_execution_context(_context(source_ref)):
        return capability.execute(capability.validate(arguments))


def _capability(rag, knowledge, provider, attachments, clock=None):
    return WeeklyReportCapability(
        rag,
        knowledge,
        provider,
        attachments,
        timezone_name="Asia/Shanghai",
        clock=clock or (lambda: datetime(2026, 10, 7, 15, 30, tzinfo=SHANGHAI)),
    )


def test_no_template_reports_the_exact_message_and_writes_nothing():
    attachments = FakeAttachments()
    capability = _capability(
        FakeRAG(),
        FakeKnowledge(attachments=[]),
        FakeProvider({}),
        attachments,
    )

    result = _run(capability, {"person": "张三", "instruction": "给张三写一份周报"})

    assert result["status"] == "no_template"
    assert result["message"] == NO_TEMPLATE_MESSAGE
    assert result["attachment_id"] == ""
    assert attachments.saved == []


def test_collected_template_produces_a_docx_and_a_source_list():
    collected = [
        {
            "attachment_id": "att-template",
            "file_name": "周报模板.docx",
            "mime_type": DOCX,
            "size_bytes": 2048,
            "content_hash": "hash-1",
            "created_at": "2026-10-06T04:00:00+00:00",
        }
    ]
    provider = FakeProvider(
        {
            "table.一_本周工作完成情况.0": ["1", "完成官网部署", "已完成", "上线 1 个站点", "✅"],
            "table.一_本周工作完成情况.1": ["2", "孩子生病请假", "已完成", "无", "✅"],
            "table.二_本周遇到的问题与风险.0": ["部署窗口冲突", "中", "已协调运维", "需要云资源"],
            "bullet.四_思考与建议_可选.0": "我建议把部署脚本固化下来",
        }
    )
    knowledge = FakeKnowledge(attachments=collected)
    attachments = FakeAttachments()
    capability = _capability(
        FakeRAG(chunks=[("r1", "message", "消息", "张三", "chunk-1", "官网部署已完成")]),
        knowledge,
        provider,
        attachments,
    )

    result = _run(
        capability,
        {"person": "张三", "instruction": "给张三写一份周报"},
    )

    assert result["status"] == "ready"
    assert result["attachment_id"] == "artifact-1"
    assert result["file_name"] == "张三2026-09-28-2026-10-04.docx"
    assert result["period"] == "上一周周报（2026-09-28 至 2026-10-04）"
    assert knowledge.opened == ["att-template"]
    assert result["citations"][0]["evidence_id"] == "chunk-1"

    # The window limits the retrieval to the previous natural week.
    bodies = capability.rag.export_bodies or capability.rag.source_bodies
    assert bodies
    for candidate in bodies:
        assert candidate["occurred_after"] == "2026-09-27T16:00:00Z"
        assert candidate["occurred_before"] == "2026-10-04T16:00:00Z"

    rendered = Document(BytesIO(attachments.saved[0]["data"]))
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "**周期：** 上一周周报（2026-09-28 至 2026-10-04）" in text
    assert "**汇报人：** 张三" in text
    assert "完成官网部署" in text
    # Private life never reaches the document.
    assert "孩子" not in text
    # The template's own title is kept.
    assert "## 周报模板（通用版）" in text


def test_uploaded_template_is_used_when_the_user_says_so():
    uploaded = {
        "up-1": {
            "attachment_id": "up-1",
            "file_name": "我的模板.docx",
            "mime_type": DOCX,
            "size_bytes": 2048,
            "minio_object_name": "u/up-1.docx",
        }
    }
    knowledge = FakeKnowledge(attachments=[])
    attachments = FakeAttachments(uploaded=uploaded)
    capability = _capability(FakeRAG(), knowledge, FakeProvider({}), attachments)

    result = _run(
        capability,
        {
            "person": "张三",
            "instruction": "按我上传的模板给张三写周报",
            "attachment_ids": ["up-1"],
        },
    )

    assert result["status"] == "ready"
    assert result["template_origin"] == "uploaded"
    assert knowledge.opened == []


def test_multiple_templates_return_a_selectable_candidate_list():
    collected = [
        {
            "attachment_id": "a1",
            "file_name": "周报模板.docx",
            "mime_type": DOCX,
            "size_bytes": 2048,
            "content_hash": "h1",
            "created_at": "2026-10-06T04:00:00+00:00",
        },
        {
            "attachment_id": "a2",
            "file_name": "周报模板.docx",
            "mime_type": DOCX,
            "size_bytes": 4096,
            "content_hash": "h2",
            "created_at": "2026-10-05T04:00:00+00:00",
        },
    ]
    attachments = FakeAttachments()
    capability = _capability(
        FakeRAG(),
        FakeKnowledge(attachments=collected),
        FakeProvider({}),
        attachments,
    )

    result = _run(capability, {"person": "张三", "instruction": "给张三写一份周报"})

    assert result["status"] == "needs_selection"
    assert [item["attachment_id"] for item in result["candidates"]] == ["a1", "a2"]
    # Each candidate is copied into the Agent store so it can be previewed.
    assert [item["preview_attachment_id"] for item in result["candidates"]] == [
        "artifact-1",
        "artifact-2",
    ]


def test_template_selection_keeps_an_explicit_document_name():
    collected = [
        {
            "attachment_id": "a1",
            "file_name": "周报模板.docx",
            "mime_type": DOCX,
            "size_bytes": 2048,
            "content_hash": "h1",
            "created_at": "2026-10-06T04:00:00+00:00",
        },
        {
            "attachment_id": "a2",
            "file_name": "周报模板.docx",
            "mime_type": DOCX,
            "size_bytes": 4096,
            "content_hash": "h2",
            "created_at": "2026-10-05T04:00:00+00:00",
        },
    ]
    capability = _capability(
        FakeRAG(),
        FakeKnowledge(attachments=collected),
        FakeProvider({}),
        FakeAttachments(),
    )

    result = _run(
        capability,
        {"person": "张三", "instruction": "给张三写周报，文件名叫季度周报"},
    )

    assert result["status"] == "needs_selection"
    assert result["title"] == "季度周报"
    assert result["file_name"] == "季度周报"


def test_chosen_template_attachment_short_circuits_the_search():
    uploaded = {
        "up-9": {
            "attachment_id": "up-9",
            "file_name": "选中的模板.docx",
            "mime_type": DOCX,
            "size_bytes": 2048,
            "minio_object_name": "u/up-9.docx",
        }
    }
    knowledge = FakeKnowledge(attachments=[])
    capability = _capability(
        FakeRAG(),
        knowledge,
        FakeProvider({}),
        FakeAttachments(uploaded=uploaded),
    )

    result = _run(
        capability,
        {"person": "张三", "instruction": "用这份模板写周报", "template_attachment_id": "up-9"},
    )

    assert result["status"] == "ready"
    assert result["template_origin"] == "uploaded"


def test_person_without_matches_reports_no_person():
    capability = _capability(
        FakeRAG(),
        FakeKnowledge(matches=[]),
        FakeProvider({}),
        FakeAttachments(),
    )

    result = _run(capability, {"person": "不存在的人", "instruction": "写周报"})

    assert result["status"] == "no_person"
    assert result["attachment_id"] == ""


def test_result_preview_carries_the_card_fields_but_not_the_document():
    from app.kernel.runtime import _result_preview

    preview = _result_preview(
        CAPABILITY_NAME,
        {
            "status": "ready",
            "message": "已按模板生成张三的周报（2026-09-28 至 2026-10-04）。",
            "title": "张三的周报",
            "subject_name": "张三",
            "period": "上一周周报（2026-09-28 至 2026-10-04）",
            "file_name": "张三2026-09-28-2026-10-04.docx",
            "attachment_id": "artifact-1",
            "size_bytes": 2048,
            "expires_at": "2026-10-08T05:00:00+00:00",
            "template_origin": "collected",
            "citations": [{"evidence_id": "c1", "title": "消息"}],
            "candidates": [],
            "data": b"never-serialised",
        },
    )

    assert preview["block_type"] == "report"
    assert preview["attachment_id"] == "artifact-1"
    assert preview["status"] == "ready"
    assert preview["title"] == "张三的周报"
    assert preview["citations"] == [{"evidence_id": "c1", "title": "消息"}]
    assert "data" not in preview


def test_provider_envelopes_are_not_shown_as_source_text():
    from app.capabilities.report import _citations

    citations = _citations(
        [
            {
                "evidence_id": "c1",
                "resource_type": "message",
                "resource_id": "m1",
                "conversation_id": "chat-1",
                "sender_name": "张三",
                "conversation_name": "私聊",
                "quote": '<msg><emoji fromusername="wxid_a" type="2" /></msg>',
            },
            {
                "evidence_id": "c2",
                "resource_type": "message",
                "resource_id": "m2",
                "conversation_id": "chat-1",
                "sender_name": "张三",
                "conversation_name": "私聊",
                "quote": "官网部署已完成",
            },
        ]
    )

    assert citations[0]["quote"] == ""
    assert citations[0]["conversation_name"] == "私聊"
    assert citations[1]["quote"] == "官网部署已完成"


def test_sources_are_one_document_per_conversation_and_one_per_message():
    from app.capabilities.report import _citations

    def document(chunk: str, quote: str, conversation: str, attachment: str) -> dict:
        return {
            "evidence_id": chunk,
            "resource_type": "attachment",
            "resource_id": attachment,
            "title": "公司简介.docx",
            "conversation_id": conversation,
            "conversation_name": conversation,
            "sender_name": "稻成",
            "quote": quote,
        }

    citations = _citations(
        [
            # One document, three chunks, one conversation -> one source.
            document("c1", "运营策略：坚持开源路线", "先躺会再说", "a1"),
            document("c2", "市场表现：2025 年 1 月 App 上线", "先躺会再说", "a1"),
            document("c3", "最新估值：估值达到 750 亿美元", "先躺会再说", "a1"),
            # The same document sent to another chat is listed separately.
            document("c4", "运营策略：坚持开源路线", "Castorice", "a2"),
            # A long message split into two chunks is still one source.
            {
                "evidence_id": "m1",
                "resource_type": "message",
                "resource_id": "msg-1",
                "conversation_id": "先躺会再说",
                "conversation_name": "先躺会再说",
                "sender_name": "先躺会再说",
                "quote": "第一段",
            },
            {
                "evidence_id": "m2",
                "resource_type": "message",
                "resource_id": "msg-1",
                "conversation_id": "先躺会再说",
                "conversation_name": "先躺会再说",
                "sender_name": "先躺会再说",
                "quote": "第二段",
            },
        ]
    )

    assert [item["title"] for item in citations] == [
        "公司简介.docx",
        "公司简介.docx",
        "",
    ]
    assert [item["conversation_name"] for item in citations] == [
        "先躺会再说",
        "Castorice",
        "先躺会再说",
    ]
    assert citations[0]["resource_type"] == "attachment"
    assert citations[2]["resource_type"] == "message"


def test_self_reference_anchors_mentions_on_the_resolved_name():
    """'给我写周报' must not search for the pronoun 我 as a mention."""

    collected = [
        {
            "attachment_id": "att-template",
            "file_name": "周报模板.docx",
            "mime_type": DOCX,
            "size_bytes": 2048,
            "content_hash": "hash-1",
            "created_at": "2026-10-06T04:00:00+00:00",
        }
    ]
    rag = FakeRAG(chunks=[("r1", "message", "消息", "稻成", "chunk-1", "开发支付模块")])
    knowledge = FakeKnowledge(
        subject="我",
        matches=[
            {
                "person_key": "u1",
                "display_name": "稻成",
                "identity_ids": ["id-1"],
                "private_conversation_ids": [],
                "identities": [
                    {"id": "id-1", "platform": "wechat", "display_name": "稻成"},
                    {"id": "id-2", "platform": "feishu", "display_name": "DC"},
                ],
            }
        ],
        attachments=collected,
    )
    attachments = FakeAttachments()
    capability = _capability(rag, knowledge, FakeProvider({}), attachments)

    result = _run(capability, {"person": "我", "instruction": "给我写周报"})

    assert result["status"] == "ready"
    assert result["subject_name"] == "DC"
    assert result["title"] == "DC的周报"
    assert result["file_name"].startswith("DC")
    # One mention query per platform name: content_contains is an AND, and the
    # pronoun 我 would match almost every sentence.
    anchors = [
        tuple(body["content_contains"])
        for body in rag.export_bodies
        if body.get("content_contains")
    ]
    assert anchors, "the mention source should still be searched"
    assert ("稻成",) in anchors
    assert ("DC",) in anchors
    assert ("我",) not in anchors


def _template_attachment() -> list[dict]:
    return [
        {
            "attachment_id": "att-template",
            "file_name": "周报模板.docx",
            "mime_type": DOCX,
            "size_bytes": 2048,
            "content_hash": "hash-1",
            "created_at": "2026-10-06T04:00:00+00:00",
        }
    ]


def test_a_self_report_is_addressed_with_the_feishu_name():
    """A self report uses the owner's own platform name, Feishu first."""

    rag = FakeRAG(chunks=[("r1", "message", "消息", "稻成", "chunk-1", "开发支付模块")])
    knowledge = FakeKnowledge(
        subject="我",
        matches=[
            {
                "person_key": "u1",
                "display_name": "wxid_v0o8tmxy91ea22",
                "identity_ids": ["id-1"],
                "private_conversation_ids": [],
                "identities": [
                    {"id": "id-1", "platform": "wechat", "display_name": "稻成"},
                    {"id": "id-2", "platform": "feishu", "display_name": "DC"},
                ],
            }
        ],
        attachments=_template_attachment(),
    )
    attachments = FakeAttachments()
    capability = _capability(rag, knowledge, FakeProvider({}), attachments)

    result = _run(
        capability,
        {"person": "我", "instruction": "帮我写一份本周周报"},
        source_ref={"owner_name": "123"},
    )

    assert result["status"] == "ready"
    assert result["subject_name"] == "DC"
    assert result["title"] == "DC的周报"
    assert result["file_name"] == "DC2026-10-05-2026-10-11.docx"
    assert "DC" in result["message"]
    # Account labels are never used as a mention anchor.
    anchors = [
        tuple(body["content_contains"])
        for body in rag.export_bodies
        if body.get("content_contains")
    ]
    assert anchors
    for value in anchors:
        assert value in {("稻成",), ("DC",)}


def test_a_self_report_without_a_profile_name_falls_back_to_the_identity():
    rag = FakeRAG(chunks=[("r1", "message", "消息", "稻成", "chunk-1", "开发支付模块")])
    knowledge = FakeKnowledge(
        subject="我",
        matches=[
            {
                "person_key": "u1",
                "display_name": "稻成",
                "identity_ids": ["id-1"],
                "private_conversation_ids": [],
            }
        ],
        attachments=_template_attachment(),
    )
    capability = _capability(rag, knowledge, FakeProvider({}), FakeAttachments())

    result = _run(capability, {"person": "我", "instruction": "帮我写周报"})

    assert result["status"] == "ready"
    assert result["subject_name"] == "稻成"
    assert result["title"] == "稻成的周报"
    assert result["file_name"].startswith("稻成")


def test_a_named_report_uses_the_name_the_user_wrote():
    """The user asked for 小李; the file must not be renamed to the contact row."""

    rag = FakeRAG(chunks=[("r1", "message", "消息", "小李", "chunk-1", "部署官网")])
    knowledge = FakeKnowledge(
        subject="小李",
        matches=[
            {
                "person_key": "p9",
                "display_name": "李四（后端）",
                "identity_ids": ["id-9"],
                "private_conversation_ids": ["conv-9"],
            }
        ],
        attachments=_template_attachment(),
    )
    capability = _capability(rag, knowledge, FakeProvider({}), FakeAttachments())

    result = _run(
        capability,
        {"person": "小李", "instruction": "给小李写一份本周周报"},
        source_ref={"owner_name": "123"},
    )

    assert result["status"] == "ready"
    assert result["title"] == "小李的周报"
    assert result["subject_name"] == "小李"
    assert result["file_name"].startswith("小李")


def test_an_explicit_document_name_overrides_the_default():
    rag = FakeRAG(chunks=[("r1", "message", "消息", "稻成", "chunk-1", "开发支付模块")])
    knowledge = FakeKnowledge(
        subject="我",
        matches=[
            {
                "person_key": "u1",
                "display_name": "稻成",
                "identity_ids": ["id-1"],
                "private_conversation_ids": [],
            }
        ],
        attachments=_template_attachment(),
    )
    capability = _capability(rag, knowledge, FakeProvider({}), FakeAttachments())

    result = _run(
        capability,
        {"person": "我", "instruction": "帮我写一份周报，文件名叫我的季度周报"},
        source_ref={"owner_name": "123"},
    )

    assert result["status"] == "ready"
    assert result["subject_name"] == "123"
    assert result["title"] == "我的季度周报"
    assert result["file_name"] == "我的季度周报.docx"


def test_an_uploaded_file_is_not_mistaken_for_an_explicit_document_name():
    uploaded = {
        "up-1": {
            "attachment_id": "up-1",
            "file_name": "我的模板.docx",
            "mime_type": DOCX,
            "size_bytes": 2048,
            "minio_object_name": "u/up-1.docx",
        }
    }
    knowledge = FakeKnowledge(attachments=[])
    capability = _capability(
        FakeRAG(),
        knowledge,
        FakeProvider({}),
        FakeAttachments(uploaded=uploaded),
    )

    result = _run(
        capability,
        {
            "person": "张三",
            "instruction": "按我上传的文件给张三写周报",
            "attachment_ids": ["up-1"],
        },
    )

    assert result["status"] == "ready"
    assert result["title"] == "张三的周报"
    assert result["file_name"].startswith("张三")


def test_a_week_before_last_report_covers_the_named_window():
    rag = FakeRAG(chunks=[("r1", "message", "消息", "张三", "chunk-1", "部署官网")])
    knowledge = FakeKnowledge(attachments=_template_attachment())
    attachments = FakeAttachments()
    capability = _capability(rag, knowledge, FakeProvider({}), attachments)

    result = _run(
        capability,
        {"person": "张三", "instruction": "生成张三的上上一周的周报"},
    )

    assert result["status"] == "ready"
    assert result["file_name"] == "张三2026-09-21-2026-09-27.docx"
    assert result["period"] == "上上周周报（2026-09-21 至 2026-09-27）"

    # The retrieval window is the week before last, not last week.
    bodies = capability.rag.export_bodies or capability.rag.source_bodies
    assert bodies
    for candidate in bodies:
        assert candidate["occurred_after"] == "2026-09-20T16:00:00Z"
        assert candidate["occurred_before"] == "2026-09-27T16:00:00Z"
